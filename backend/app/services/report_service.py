"""PDF and Excel exports of a course's attainment report.

Both formats render the same data as `attainment_service.build_course_report`
returns, so the numbers on a downloaded report always match the dashboard.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.course import Course
from app.schemas.attainment import CourseAttainmentReport

ACHIEVED_FILL = "C6EFCE"
NOT_ACHIEVED_FILL = "FFC7CE"


def _clo_table_rows(report: CourseAttainmentReport) -> list[list[str]]:
    header = ["CLO Code", "Title", "Class Average (%)", "Students", "Achieved"]
    rows = [header]
    for clo in report.clo_attainment:
        rows.append(
            [
                clo.code,
                clo.title,
                f"{clo.class_average:.2f}",
                str(clo.student_count),
                "Yes" if clo.is_achieved else "No",
            ]
        )
    return rows


def _plo_table_rows(report: CourseAttainmentReport) -> list[list[str]]:
    header = ["PLO Code", "Title", "Class Average (%)"]
    rows = [header]
    for plo in report.plo_attainment:
        rows.append([plo.code, plo.title, f"{plo.class_average:.2f}"])
    return rows


def build_pdf_report(course: Course, report: CourseAttainmentReport) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter, topMargin=0.75 * inch, bottomMargin=0.75 * inch
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"Attainment Report: {course.code} — {course.name}", styles["Title"]),
        Paragraph(
            f"Threshold: {report.threshold:.2f}% &nbsp;&nbsp; "
            f"Enrolled students: {report.student_count}",
            styles["Normal"],
        ),
        Spacer(1, 0.3 * inch),
        Paragraph("CLO Attainment", styles["Heading2"]),
        _styled_table(_clo_table_rows(report), achieved_col=4),
        Spacer(1, 0.3 * inch),
        Paragraph("PLO Attainment", styles["Heading2"]),
        _styled_table(_plo_table_rows(report))
        if report.plo_attainment
        else Paragraph("No PLO mappings confirmed for this course yet.", styles["Normal"]),
    ]

    doc.build(story)
    return buffer.getvalue()


def _styled_table(rows: list[list[str]], achieved_col: int | None = None) -> Table:
    table = Table(rows, hAlign="LEFT")
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]

    if achieved_col is not None:
        for row_index, row in enumerate(rows[1:], start=1):
            fill = colors.HexColor(
                "#dcfce7" if row[achieved_col] == "Yes" else "#fee2e2"
            )
            style.append(("BACKGROUND", (achieved_col, row_index), (achieved_col, row_index), fill))

    table.setStyle(TableStyle(style))
    return table


def build_excel_report(course: Course, report: CourseAttainmentReport) -> bytes:
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = f"Attainment Report: {course.code} — {course.name}"
    summary["A1"].font = Font(size=14, bold=True)
    summary["A2"] = "Threshold (%)"
    summary["B2"] = report.threshold
    summary["A3"] = "Enrolled students"
    summary["B3"] = report.student_count

    _write_sheet(
        workbook.create_sheet("CLO Attainment"),
        _clo_table_rows(report),
        achieved_col=4,
    )
    _write_sheet(workbook.create_sheet("PLO Attainment"), _plo_table_rows(report))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_sheet(sheet, rows: list[list[str]], achieved_col: int | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="1E293B")
    header_font = Font(color="FFFFFF", bold=True)

    for row in rows:
        sheet.append(row)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    if achieved_col is not None:
        for row_index in range(2, len(rows) + 1):
            cell = sheet.cell(row=row_index, column=achieved_col + 1)
            achieved = cell.value == "Yes"
            cell.fill = PatternFill(
                "solid", fgColor=ACHIEVED_FILL if achieved else NOT_ACHIEVED_FILL
            )

    for column_index, column_cells in enumerate(sheet.columns, start=1):
        longest = max((len(str(cell.value)) for cell in column_cells if cell.value), default=10)
        sheet.column_dimensions[get_column_letter(column_index)].width = min(longest + 4, 50)
