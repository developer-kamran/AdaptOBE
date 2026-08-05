"""PDF and Excel export endpoints."""

from openpyxl import load_workbook

from tests.conftest import auth_header


async def test_export_pdf_requires_course_access(client, make_user, course):
    other = await make_user("faculty.export@adaptobe.edu", role="faculty")
    resp = await client.get(
        f"/api/v1/attainment/course/{course.id}/export/pdf", headers=auth_header(other)
    )
    assert resp.status_code == 403


async def test_export_pdf_returns_pdf_bytes(client, faculty, course, make_clo):
    await make_clo("CLO-1", "Loops", "Write loops and conditionals.")
    resp = await client.get(
        f"/api/v1/attainment/course/{course.id}/export/pdf", headers=auth_header(faculty)
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF")
    assert course.code in resp.headers["content-disposition"]


async def test_export_excel_returns_valid_workbook(
    client, db_session, faculty, course, make_clo, make_plo, make_student, enroll,
    make_assessment, make_question, enter_scores,
):
    import io

    clo = await make_clo("CLO-1", "Loops", "Write loops and conditionals.")
    plo = await make_plo("PLO-1", "Knowledge", "Apply engineering knowledge.")

    from app.schemas.mapping import MappingConfirmRequest
    from app.services import mapping_service

    await mapping_service.confirm_mapping(
        db_session, MappingConfirmRequest(clo_id=clo.id, plo_id=plo.id, strength=2), faculty
    )

    student = await make_student("excel")
    await enroll(student)
    assessment = await make_assessment(total_marks=10.0)
    question = await make_question(assessment.id, 1, 10.0, clo.id)
    await enter_scores(assessment.id, [(question.id, student.id, 8.0)])

    resp = await client.get(
        f"/api/v1/attainment/course/{course.id}/export/excel", headers=auth_header(faculty)
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(io.BytesIO(resp.content))
    assert workbook.sheetnames == ["Summary", "CLO Attainment", "PLO Attainment"]

    clo_sheet = workbook["CLO Attainment"]
    assert clo_sheet["A1"].value == "CLO Code"
    assert clo_sheet["A2"].value == "CLO-1"
    assert clo_sheet["C2"].value == "80.00"

    plo_sheet = workbook["PLO Attainment"]
    assert plo_sheet["A2"].value == "PLO-1"
    assert plo_sheet["C2"].value == "80.00"


async def test_export_pdf_with_no_plo_mappings_does_not_error(client, faculty, course, make_clo):
    """A course with CLOs but no confirmed PLO mappings should still export cleanly."""
    await make_clo("CLO-1", "Loops", "Write loops and conditionals.")
    resp = await client.get(
        f"/api/v1/attainment/course/{course.id}/export/pdf", headers=auth_header(faculty)
    )
    assert resp.status_code == 200
    assert resp.content.startswith(b"%PDF")
