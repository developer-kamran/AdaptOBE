"""Per-assessment exam-paper PDF/Word export."""

import io

import docx
import pdfplumber

from tests.conftest import auth_header


async def test_export_pdf_requires_faculty_owner(client, make_user, make_assessment):
    other = await make_user("faculty.aexport@adaptobe.edu", role="faculty", employee_id="FAC-AEXPORT")
    assessment = await make_assessment(total_marks=10.0)
    resp = await client.get(
        f"/api/v1/assessments/{assessment.id}/export/pdf", headers=auth_header(other)
    )
    assert resp.status_code == 403


async def test_export_pdf_returns_valid_pdf_with_content(
    client, faculty, course, make_assessment, make_question
):
    assessment = await make_assessment(
        title="Midterm Exam", total_marks=10.0, weightage_percent=20.0
    )
    await make_question(assessment.id, 1, 10.0, text="Explain process scheduling.")

    resp = await client.get(
        f"/api/v1/assessments/{assessment.id}/export/pdf", headers=auth_header(faculty)
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF")
    assert "inline" in resp.headers["content-disposition"]

    with pdfplumber.open(io.BytesIO(resp.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert course.code in text
    assert "Midterm Exam" in text
    assert "Explain process scheduling." in text
    assert faculty.full_name in text


async def test_export_docx_returns_valid_document_with_content(
    client, faculty, course, make_assessment, make_question
):
    assessment = await make_assessment(title="Quiz 1", total_marks=5.0, weightage_percent=5.0)
    await make_question(assessment.id, 1, 5.0, text="Define a process.")

    resp = await client.get(
        f"/api/v1/assessments/{assessment.id}/export/docx", headers=auth_header(faculty)
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert resp.content.startswith(b"PK")
    assert "attachment" in resp.headers["content-disposition"]

    document = docx.Document(io.BytesIO(resp.content))
    # The Total Marks/Course Code/Instructor/Allocated Time/Date block lives in
    # a table (left column: marks/code/instructor, right: time/date), not a
    # top-level paragraph -- `document.paragraphs` doesn't see inside tables.
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    text = "\n".join(p.text for p in document.paragraphs) + "\n" + table_text
    assert course.code in text
    assert "Define a process." in text


async def test_export_pdf_handles_assessment_with_no_questions(client, faculty, make_assessment):
    assessment = await make_assessment(total_marks=10.0)
    resp = await client.get(
        f"/api/v1/assessments/{assessment.id}/export/pdf", headers=auth_header(faculty)
    )
    assert resp.status_code == 200
    assert resp.content.startswith(b"%PDF")


async def test_export_pdf_shows_allocated_time_when_set(
    client, db_session, faculty, make_assessment
):
    assessment = await make_assessment(total_marks=10.0)
    assessment.duration_minutes = 90
    await db_session.commit()

    resp = await client.get(
        f"/api/v1/assessments/{assessment.id}/export/pdf", headers=auth_header(faculty)
    )
    with pdfplumber.open(io.BytesIO(resp.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "90 minutes" in text
