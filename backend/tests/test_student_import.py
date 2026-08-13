"""Bulk student import: parsing, AI column matching, preview, and commit.

Note on the encoder: `conftest.stub_encoder` swaps the real sentence model
for a bag-of-words fake so the suite never loads PyTorch. That fake ranks
differently from `all-MiniLM-L6-v2`, so the tests below exercise the
*matching logic* (alias fast path, threshold, greedy assignment) with
controlled vectors rather than asserting on the real model's semantic
judgement, which is model behaviour and not our code.
"""

from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.core.security import generate_password
from app.ml import column_matcher, embeddings
from app.models.user import User, UserRole
from app.services import file_parsers
from tests.conftest import auth_header

HEADERS = ["Full Name", "Father Name", "Enrollment No", "Seat No", "Email"]


def build_xlsx(rows: list[list[str]], headers: list[str] | None = None) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers if headers is not None else HEADERS)
    for row in rows:
        worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def upload(content: bytes, filename: str = "roster.xlsx"):
    return {
        "file": (
            filename,
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }


@pytest.fixture(autouse=True)
def _reset_prompt_cache():
    """`column_matcher` caches its prompt embeddings module-wide; clear it so
    a test that swaps the encoder isn't served vectors from another test."""
    column_matcher._prompt_embeddings_cache.clear()
    yield
    column_matcher._prompt_embeddings_cache.clear()


# --- Pure parsing ------------------------------------------------------------


def test_parse_excel_reads_header_and_rows():
    content = build_xlsx([["Ali Raza", "Nasir Raza", "ENR-1", "SEAT-1", "ali@x.edu"]])
    rows = file_parsers.parse_excel(content)

    assert rows[0] == HEADERS
    assert rows[1] == ["Ali Raza", "Nasir Raza", "ENR-1", "SEAT-1", "ali@x.edu"]


def test_parse_excel_stringifies_numeric_cells_without_decimals():
    """An enrollment number typed as a number must not become '12345.0'."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(HEADERS)
    worksheet.append(["Ali", "Nasir", 12345, 67890, "ali@x.edu"])
    buffer = BytesIO()
    workbook.save(buffer)

    rows = file_parsers.parse_excel(buffer.getvalue())
    assert rows[1][2] == "12345"
    assert rows[1][3] == "67890"


def test_parse_excel_skips_blank_rows():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(HEADERS)
    worksheet.append([None, None, None, None, None])
    worksheet.append(["Ali", "Nasir", "ENR-1", "SEAT-1", "ali@x.edu"])
    buffer = BytesIO()
    workbook.save(buffer)

    rows = file_parsers.parse_excel(buffer.getvalue())
    assert len(rows) == 2  # header + the one real row


def test_parse_file_rejects_unsupported_extension():
    with pytest.raises(file_parsers.FileParseError):
        file_parsers.parse_file(b"whatever", "roster.txt")


def test_parse_excel_raises_on_corrupt_content():
    with pytest.raises(file_parsers.FileParseError):
        file_parsers.parse_excel(b"this is definitely not a workbook")


# --- Column matching ---------------------------------------------------------


async def test_match_columns_uses_alias_fast_path():
    """Obvious headers resolve exactly, with no embedding involved."""
    detected = await column_matcher.match_columns(HEADERS)

    assert detected == {
        "full_name": 0,
        "father_name": 1,
        "enrollment_no": 2,
        "seat_no": 3,
        "email": 4,
    }


async def test_match_columns_alias_is_case_and_punctuation_insensitive():
    detected = await column_matcher.match_columns(["FULL_NAME", "father's name", "Roll No."])

    assert detected["full_name"] == 0
    assert detected["father_name"] == 1
    assert detected["seat_no"] == 2


async def test_match_columns_falls_back_to_embeddings(monkeypatch):
    """A header no alias covers is matched by meaning instead."""
    guardian = [1.0] + [0.0] * (embeddings.EMBEDDING_DIM - 1)

    def encoder(text: str) -> list[float]:
        # "Ward" is unlike any alias; make it point at the father_name prompt.
        if text in ("Ward", column_matcher.FIELD_PROMPTS["father_name"]):
            return guardian
        return [0.0, 1.0] + [0.0] * (embeddings.EMBEDDING_DIM - 2)

    monkeypatch.setattr(embeddings, "encode_text", encoder)

    detected = await column_matcher.match_columns(["Ward"])
    assert detected["father_name"] == 0


async def test_match_columns_ignores_unrelated_headers(monkeypatch):
    """A column below the similarity threshold matches nothing."""

    def encoder(text: str) -> list[float]:
        if text == "CGPA":
            return [0.0, 1.0] + [0.0] * (embeddings.EMBEDDING_DIM - 2)
        return [1.0] + [0.0] * (embeddings.EMBEDDING_DIM - 1)

    monkeypatch.setattr(embeddings, "encode_text", encoder)

    detected = await column_matcher.match_columns(["CGPA"])
    assert all(index is None for index in detected.values())


async def test_match_columns_never_assigns_one_column_to_two_fields(monkeypatch):
    """Identical vectors for every prompt must still yield a 1:1 mapping."""
    same = [1.0] + [0.0] * (embeddings.EMBEDDING_DIM - 1)
    monkeypatch.setattr(embeddings, "encode_text", lambda text: same)

    detected = await column_matcher.match_columns(["Mystery A", "Mystery B"])

    assigned = [index for index in detected.values() if index is not None]
    assert len(assigned) == len(set(assigned))


def test_cosine_similarity_of_identical_unit_vectors_is_one():
    vector = embeddings.encode_text("father name")
    assert column_matcher.cosine_similarity(vector, vector) == pytest.approx(1.0)


# --- Password generation -----------------------------------------------------


def test_generate_password_is_eight_chars_with_letter_and_digit():
    for _ in range(50):
        password = generate_password()
        assert len(password) == 8
        assert any(c.isalpha() for c in password)
        assert any(c.isdigit() for c in password)
        # Ambiguous glyphs would be misread off a printed handout.
        assert not set(password) & set("0O1lI")


def test_generate_password_values_differ():
    assert len({generate_password() for _ in range(50)}) > 40


# --- Preview endpoint --------------------------------------------------------


async def test_preview_counts_complete_and_incomplete(client, sub_admin):
    content = build_xlsx(
        [
            ["Ali Raza", "Nasir Raza", "ENR-1", "SEAT-1", "ali@x.edu"],
            ["Sara Khan", "", "ENR-2", "SEAT-2", "sara@x.edu"],  # no father's name
        ]
    )

    resp = await client.post(
        "/api/v1/admin/students/import/preview",
        files=upload(content),
        headers=auth_header(sub_admin),
    )
    assert resp.status_code == 200
    body = resp.json()

    assert body["total_detected"] == 2
    assert body["complete_count"] == 1
    assert body["incomplete_count"] == 1

    incomplete = next(r for r in body["records"] if not r["is_complete"])
    assert incomplete["missing_fields"] == ["Father's Name"]
    assert body["detected_columns"]["father_name"] == "Father Name"


async def test_preview_creates_nothing(client, db_session, sub_admin):
    content = build_xlsx([["Ali Raza", "Nasir Raza", "ENR-1", "SEAT-1", "ali@x.edu"]])

    resp = await client.post(
        "/api/v1/admin/students/import/preview",
        files=upload(content),
        headers=auth_header(sub_admin),
    )
    assert resp.status_code == 200

    found = await db_session.execute(select(User).where(User.email == "ali@x.edu"))
    assert found.scalar_one_or_none() is None


async def test_preview_flags_invalid_email_and_duplicates_within_file(client, sub_admin):
    content = build_xlsx(
        [
            ["Ali Raza", "Nasir Raza", "ENR-1", "SEAT-1", "not-an-email"],
            ["Sara Khan", "Imran Khan", "ENR-2", "SEAT-2", "dup@x.edu"],
            ["Bilal Ahmed", "Nasir Ahmed", "ENR-3", "SEAT-3", "dup@x.edu"],
        ]
    )

    resp = await client.post(
        "/api/v1/admin/students/import/preview",
        files=upload(content),
        headers=auth_header(sub_admin),
    )
    body = resp.json()

    assert body["complete_count"] == 1  # only the first 'dup@x.edu' row survives
    assert "Email address is not valid" in body["records"][0]["issues"]
    assert "Duplicate email in this file" in body["records"][2]["issues"]


async def test_preview_flags_duplicate_seat_number_within_file(client, sub_admin):
    """Enrollment No and Seat No must each be unique -- a clash on either one
    should be flagged in-file, not just discovered at confirm time."""
    content = build_xlsx(
        [
            ["Ali Raza", "Nasir Raza", "ENR-1", "SEAT-DUP", "ali@x.edu"],
            ["Sara Khan", "Imran Khan", "ENR-2", "SEAT-DUP", "sara@x.edu"],
        ]
    )

    resp = await client.post(
        "/api/v1/admin/students/import/preview",
        files=upload(content),
        headers=auth_header(sub_admin),
    )
    body = resp.json()

    assert body["complete_count"] == 1  # only the first 'SEAT-DUP' row survives
    assert body["records"][0]["is_complete"] is True
    assert "Duplicate seat number in this file" in body["records"][1]["issues"]


async def test_preview_warns_when_a_required_column_is_absent(client, sub_admin):
    content = build_xlsx(
        [["Ali Raza", "ENR-1", "SEAT-1", "ali@x.edu"]],
        headers=["Full Name", "Enrollment No", "Seat No", "Email"],
    )

    resp = await client.post(
        "/api/v1/admin/students/import/preview",
        files=upload(content),
        headers=auth_header(sub_admin),
    )
    body = resp.json()

    assert body["detected_columns"]["father_name"] is None
    assert body["complete_count"] == 0
    assert any("Father's Name" in warning for warning in body["warnings"])


async def test_preview_rejects_unsupported_extension(client, sub_admin):
    resp = await client.post(
        "/api/v1/admin/students/import/preview",
        files={"file": ("roster.txt", b"nope", "text/plain")},
        headers=auth_header(sub_admin),
    )
    assert resp.status_code == 422


async def test_preview_rejects_empty_file(client, sub_admin):
    resp = await client.post(
        "/api/v1/admin/students/import/preview",
        files=upload(b""),
        headers=auth_header(sub_admin),
    )
    assert resp.status_code == 422


# --- Confirm endpoint --------------------------------------------------------


def confirm_payload(*students):
    return {"students": list(students)}


def student_row(suffix: str) -> dict:
    return {
        "full_name": f"Student {suffix}",
        "father_name": f"Father {suffix}",
        "enrollment_no": f"ENR-{suffix}",
        "seat_no": f"SEAT-{suffix}",
        "email": f"student.{suffix}@x.edu",
    }


async def test_confirm_creates_students_in_callers_department(
    client, db_session, sub_admin, department
):
    resp = await client.post(
        "/api/v1/admin/students/import/confirm",
        json=confirm_payload(student_row("a"), student_row("b")),
        headers=auth_header(sub_admin),
    )
    assert resp.status_code == 200
    body = resp.json()

    assert body["created_count"] == 2
    assert body["skipped_count"] == 0
    assert all(len(r["generated_password"]) == 8 for r in body["results"])
    # Distinct password per student.
    assert len({r["generated_password"] for r in body["results"]}) == 2

    created = await db_session.execute(select(User).where(User.email == "student.a@x.edu"))
    student = created.scalar_one()
    assert student.role == UserRole.student
    assert student.dept_id == department.id  # forced to the caller's department
    assert student.father_name == "Father a"
    assert student.password_encrypted is not None


async def test_confirm_generated_password_actually_works_for_login(client, sub_admin):
    resp = await client.post(
        "/api/v1/admin/students/import/confirm",
        json=confirm_payload(student_row("login")),
        headers=auth_header(sub_admin),
    )
    password = resp.json()["results"][0]["generated_password"]

    login = await client.post(
        "/api/v1/auth/login",
        json={"email": "student.login@x.edu", "password": password},
    )
    assert login.status_code == 200


async def test_confirm_skips_duplicate_without_aborting_the_batch(client, make_user, sub_admin):
    await make_user(
        "student.taken@x.edu",
        role=UserRole.student,
        enrollment_no="ENR-taken",
        seat_no="SEAT-taken",
        father_name="Father taken",
    )

    resp = await client.post(
        "/api/v1/admin/students/import/confirm",
        json=confirm_payload(student_row("taken"), student_row("fresh")),
        headers=auth_header(sub_admin),
    )
    assert resp.status_code == 200
    body = resp.json()

    assert body["created_count"] == 1
    assert body["skipped_count"] == 1
    skipped = next(r for r in body["results"] if r["status"] == "skipped")
    assert skipped["email"] == "student.taken@x.edu"
    assert skipped["generated_password"] is None


async def test_confirm_rejects_incomplete_rows_server_side(client, sub_admin):
    """The browser claiming a row is complete is not taken on trust."""
    incomplete = student_row("bad")
    incomplete["father_name"] = ""

    resp = await client.post(
        "/api/v1/admin/students/import/confirm",
        json=confirm_payload(incomplete),
        headers=auth_header(sub_admin),
    )
    assert resp.status_code == 422


async def test_confirm_rejects_malformed_email_server_side(client, sub_admin):
    bad = student_row("bademail")
    bad["email"] = "not-an-email"

    resp = await client.post(
        "/api/v1/admin/students/import/confirm",
        json=confirm_payload(bad),
        headers=auth_header(sub_admin),
    )
    assert resp.status_code == 422


# --- RBAC --------------------------------------------------------------------


async def test_import_endpoints_reject_faculty_and_super_admin(client, faculty, super_admin):
    """Importing a roster is department-scoped work: faculty don't manage
    accounts at all, and a super_admin manages departments/sub-admins only."""
    content = build_xlsx([["Ali", "Nasir", "ENR-1", "SEAT-1", "ali@x.edu"]])

    for user in (faculty, super_admin):
        preview = await client.post(
            "/api/v1/admin/students/import/preview",
            files=upload(content),
            headers=auth_header(user),
        )
        assert preview.status_code == 403, f"{user.role} should not preview imports"

        confirm = await client.post(
            "/api/v1/admin/students/import/confirm",
            json=confirm_payload(student_row("x")),
            headers=auth_header(user),
        )
        assert confirm.status_code == 403, f"{user.role} should not confirm imports"


async def test_import_endpoints_reject_students(client, make_user):
    student = await make_user(
        "student.importer@x.edu",
        role=UserRole.student,
        enrollment_no="ENR-IMPORTER",
        seat_no="SEAT-IMPORTER",
        father_name="Father Importer",
    )
    content = build_xlsx([["Ali", "Nasir", "ENR-1", "SEAT-1", "ali@x.edu"]])

    resp = await client.post(
        "/api/v1/admin/students/import/preview",
        files=upload(content),
        headers=auth_header(student),
    )
    assert resp.status_code == 403


async def test_import_endpoints_require_authentication(client):
    content = build_xlsx([["Ali", "Nasir", "ENR-1", "SEAT-1", "ali@x.edu"]])
    resp = await client.post("/api/v1/admin/students/import/preview", files=upload(content))
    assert resp.status_code == 401
